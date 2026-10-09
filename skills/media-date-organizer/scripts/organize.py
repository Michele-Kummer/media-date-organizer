#!/usr/bin/env python3
"""
organize.py — read photo and video metadata (EXIF, MOV mvhd, HEIC TIFF) and
sort files by real Date Taken.

Phases (several can be given on one command line; they run in a fixed order
and a summary of every phase's counts and problems is printed at the end):
    --all               Shorthand for --delete-junk --move-incomplete
                        --fill-blanks --convert-png --apply-moves
                        --sync-timestamps. Add --quarantine-ads or --flatten
                        to include those too.
    --delete-junk       Permanently delete every Windows thumbnail cache
                        (Thumbs.db, ehthumbs.db, ehthumbs_vista.db; matched by
                        whole name, so photos with "thumb" in the name are
                        safe) and every .AAE file in the root folder
                        and every subfolder. AAE files are the
                        edit-instruction sidecars iPhones export next to a
                        photo; they hold no image and nothing on Windows
                        reads them. Also deletes macOS "._" sidecars
                        (._IMG_0001.png beside IMG_0001.png), which a Mac
                        writes when copying to a non-Mac drive; a file is
                        only treated as one when its contents carry the
                        AppleDouble signature as well as the name. Add
                        --dry-run to list them without deleting anything.
    --report-only       Scan and write audit XLSX; make no changes.
    --fill-blanks       Write DateTimeOriginal into EXIF for JPEG/PNG files that
                        are missing Date Taken. Uses a date embedded in the
                        filename (e.g. 2023-12-20_<id>-main.jpg,
                        IMG_20231220_142355.jpg) first, then the containing
                        folder's YYYY-MM-DD name at 12:00:00. The file's
                        modified time is never used — it is usually the import
                        time. JPEGs are rewritten without re-encoding. HEIC
                        and video files are written with exiftool when it is
                        installed, else flagged. Empty (0-byte) files are
                        skipped and counted.
                        Videos (MOV/MP4/M4V/3GP) with a blank Media Created are
                        found in the root and every subfolder at any depth.
                        Their date comes from the filename first; else the
                        nearest containing folder whose name holds a full
                        date (2019-08-30, 11-11-2019) at 12:00:00. A folder
                        name holding only a year ("2014") is not used; such
                        videos are listed as skipped. The time is
                        taken as local time and stored as UTC, so Explorer
                        shows it unchanged. Add --dry-run to list what would
                        be written without changing anything.
    --convert-png       Re-encode PNG files as JPEG at quality 95, preserving
                        EXIF. Leaves originals in place. PNGs with any
                        transparent pixels are skipped and stay PNG, as are
                        "-overlay" PNGs (Snapchat caption layers). A PNG that
                        cannot be read as a picture is reported, counted as
                        read-failed and left alone; the rest still convert.
    --apply-moves       Move files whose Date Taken does not match their current
                        folder name into a sibling folder named for the real
                        date. A file with no Date Taken or Media Created is
                        not moved, but in a YYYY-MM-DD folder it gets a
                        _<folder-date> suffix, unless its name holds a date
                        already, so the folder's date survives in the name.
                        A suffix is never added to a file that has a date
                        and never stripped from any file. Extensions are
                        left exactly as they are, case included; correcting
                        them is Fix-Extensions.ps1's job.
    --compare-collisions
                        Report only; not part of --all. For every file
                        --apply-moves skips as a COLLISION (its name is
                        already taken in the folder for its real date),
                        compare it byte for byte with the file holding that
                        name and list the pair as IDENTICAL or DIFFERENT, with
                        both sizes and dates. Also covers --flatten: a file
                        it renamed with a _<date-taken> suffix, or left in
                        its subfolder, because its name was taken in the
                        root. Nothing is moved or deleted.
    --visual-compare    Dry run only (give --dry-run with it); not part of
                        --all. Compare the pixels of each pair of possible
                        duplicates and report DUPLICATE, DIFFERENT or UNSURE
                        with a confidence level (high, medium, low) and the
                        measured difference. Pairs are the collisions above
                        plus files in one folder that share a name and differ
                        only in extension (IMG_1.jpg / IMG_1.png). Byte-
                        identical pairs are IDENTICAL with no comparison
                        needed. A photo is compared as one picture, a video
                        by four frames and its length. Runs entirely on this
                        computer; videos need ffmpeg (pip install
                        imageio-ffmpeg). The confidence level gets its own
                        column in the audit workbook's Run log. Nothing is
                        moved or deleted.
    --sync-timestamps   Set each file's "Date Created" to its EXIF Date Taken,
                        and its "Date Modified" too BUT only for files that
                        appear to be unedited (EXIF ModifyDate equals
                        DateTimeOriginal and no editor-software tag is
                        present). An edited file keeps its Date Modified and
                        is counted as created-only-edited. Date Created is
                        only settable on Windows; elsewhere just Date
                        Modified is synced, and edited files are skipped.
    --flatten           Move every file that has a Date Taken or Media Created
                        out of its subfolder and up into the root folder
                        itself, then delete the subfolders left empty. Files
                        with no Date Taken or Media Created stay where they are
                        (the folder name may be the only record of their date),
                        so their folders survive. A name already taken in the
                        root gets a _<date-taken> suffix. Exception: inside the
                        _ads folder every file moves up from its subfolder into
                        _ads itself, with or without a Date Taken (a name clash
                        gets a _<subfolder-name> suffix). Camcorder videos
                        (.m2ts, .mts) and older MPEG, Windows Media and AVI
                        videos (.mpg, .mpeg, .wmv, .avi), which no other phase handles, are moved
                        too, going by the date in the filename
                        (20180116122434), else by the date in the subfolder
                        name (2018-01-16, 7-17-2014), which is then appended
                        to the filename; ones with neither, or holding no
                        data, stay. A file's .modd/.moff sidecars
                        move with it. A .modd/.moff in the root or one of its
                        subfolders whose file is not beside it is permanently
                        deleted. Ends with an OPPOSITE TYPE line: how many
                        files of the less common kind the root itself then
                        holds (videos in a photo folder, or photos in a video
                        folder), also counted as videos-in-photo-folder or
                        photos-in-video-folder. Add --dry-run to preview
                        without changing anything.
    --opposite-type     Report only; not part of --all. Give the folder that
                        holds the year folders as --root (D:\\Videos, not
                        D:\\Videos\\2014). Lists every photo filed among the
                        videos, or video among the photos, then a count for
                        each year folder (2014, "2013 West") and a total.
                        Whichever kind the root holds more of decides what
                        it is. Each year folder is searched at every depth;
                        _ads and _incomplete are not counted. A root with no
                        year folders is reported as one folder. Nothing is
                        moved or deleted.
    --final-check       Not part of --all. Give the folder that holds the
                        year folders as --root. For each year folder,
                        reports what the other phases left behind: files
                        still in an _ads folder (IN ADS), files still in an
                        _incomplete folder (INCOMPLETE), photos among videos
                        or videos among photos (as --opposite-type lists
                        them), subfolders still there after --flatten (NOT
                        FLATTENED, with how many files each holds), name
                        collisions --flatten met (COLLISION, one line per
                        pair, saying whether the two files are identical or
                        different and whether the file was renamed with a
                        date or left in its subfolder) and audit
                        workbooks (AUDIT FILE). Ends with one line per year
                        folder and a FINAL CHECK total. A root with no year
                        folders is reported as one folder. Its one change:
                        an audit workbook not yet named for its folder and
                        media type is renamed so that it is
                        (2014/media_audit_2026-10-07.xlsx becomes
                        2014_videos_media_audit_2026-10-07.xlsx). Nothing
                        else is moved, renamed or deleted. Add --dry-run to
                        list the renames without making them.
    --quarantine-ads    Move ad images and videos into an _ads folder inside
                        the root, keeping their subfolder path, for the user to
                        review and delete. A file is an ad only if BOTH hold:
                        its name is a UUID
                        (0c4cda27-b4cc-4e92-a446-d6b780f24a64), a 9-character
                        hex id plus a number (bed51b94a_1595), gmsnet plus an
                        optional number (gmsnet2), news_images%2F plus a number
                        (news_images%2F1714658759372), UnityAdsCache- plus a
                        64-character hex hash, a 24-character hex id plus
                        parts joined by - or _
                        (54ac2fda0a6755305200011c-b30-600,
                        55d58970f6cd4574f635d6e3_568-1443241239), a name
                        holding pixel dimensions
                        and ending in ___ plus a 6-character id
                        (Update_Now_Video_V2_720x1280_15s___fudxlv), or, for
                        videos only, 20 lowercase letters and digits
                        (32129eda9b8e718c5277, vuyyzy0brvod5gcaocnv), a
                        32-character hex hash, or a name ending
                        -<width>x<height>-Q2 or -<width>x<height>-h264-Q2; AND
                        it carries no camera info (EXIF Make/Model for a photo,
                        a device-make tag for a video). Name matches that do
                        have camera info are kept and listed. Looks in the root
                        folder and every subfolder. Nothing is deleted. Add
                        --dry-run to preview without moving anything.
    --move-incomplete   Move incomplete media files into an "_incomplete"
                        folder inside the root, keeping their subfolder path
                        (so each stays in its dated folder). A file is
                        incomplete if it is 0 bytes, nothing but null bytes,
                        or a video with no video header (a stub or truncated
                        copy that no player can open): the remains of a
                        failed copy or transfer. Looks in the root folder and
                        every subfolder. Nothing is deleted. An existing
                        "_no-image-data" folder from earlier versions is
                        folded into "_incomplete". --flatten leaves the
                        folder alone. --move-no-data is the old name for this
                        flag and still works. Add --dry-run to preview
                        without moving anything.

Every run appends what it reported to the "Run log" and "Run counts" sheets
of the audit workbook, dry runs included, so the workbook is a history of
everything done to the folder. It is
<root folder name>_<photos|videos>_media_audit_<date>.xlsx in the root
(2014_videos_media_audit_2026-10-07.xlsx), <date> being the day of the run
and photos or videos whichever kind the root holds more of: runs on the
same day share a workbook, a new day starts a new one. Close it in Excel
before running, or the run cannot be recorded.

Dependencies: Pillow (`pip install pillow`) and, for the audit workbook,
openpyxl (`pip install openpyxl`). Scanning works without openpyxl — the XLSX
step is skipped with a clear message if it's missing. exiftool is optional:
when present (on PATH, or in winget's default folder) --fill-blanks uses it
for HEIC and video files. Install with `winget install OliverBetz.ExifTool`.
"""
from __future__ import annotations

import argparse
import datetime
import filecmp
import io
import json
import os
import pathlib
import re
import shutil
import struct
import subprocess
import sys
from collections import Counter
from typing import Optional

try:
    from PIL import Image, ImageChops, ImageStat
    HAS_PIL = True
except ImportError:
    HAS_PIL = False
    print("WARNING: Pillow is not installed. Install with: pip install pillow", file=sys.stderr)


# ---------------- Format detection ----------------

def detect_type(path: str) -> Optional[str]:
    """Return the canonical lowercase extension (with dot) or None."""
    try:
        with open(path, 'rb') as f:
            h = f.read(16)
    except OSError:
        return None
    if len(h) < 12:
        return None
    if h[:8] == b'\x89PNG\r\n\x1a\n':
        return '.png'
    if h[:3] == b'\xff\xd8\xff':
        return '.jpg'
    if h[:4] == b'GIF8':
        return '.gif'
    if h[:2] == b'BM':
        return '.bmp'
    if h[:4] in (b'II*\x00', b'MM\x00*'):
        return '.tif'
    if h[4:8] == b'ftyp':
        brand = h[8:12]
        if brand in (b'heic', b'heix', b'mif1', b'msf1'):
            return '.heic'
        if brand == b'qt  ':
            return '.mov'
        if brand in (b'mp42', b'mp41', b'isom'):
            return '.mp4'
        if brand.startswith(b'M4V'):
            return '.m4v'
        if brand.startswith(b'3g2'):
            return '.3g2'
        if brand.startswith(b'3g'):
            return '.3gp'
        return '.mp4'
    if h[:4] == b'RIFF' and h[8:12] == b'WEBP':
        return '.webp'
    return None


# ---------------- Date Taken extraction ----------------

def _heic_dates(path: str) -> dict:
    """Return {DateTimeOriginal, DateTimeDigitized, DateTime, Software, Make,
    Model} strings found in a HEIC/HEIF file's embedded TIFF/EXIF block."""
    out = {}
    with open(path, 'rb') as f:
        data = f.read()
    for pat, endian in ((b'II*\x00', '<'), (b'MM\x00*', '>')):
        start = 0
        while True:
            i = data.find(pat, start)
            if i < 0:
                break
            start = i + 1
            try:
                ifd0 = struct.unpack(endian + 'I', data[i + 4:i + 8])[0]
                nentries = struct.unpack(endian + 'H', data[i + ifd0:i + ifd0 + 2])[0]
                if not (0 < nentries < 200):
                    continue
                for k in range(nentries):
                    e = i + ifd0 + 2 + k * 12
                    tag, typ, cnt = struct.unpack(endian + 'HHI', data[e:e + 8])
                    val_off = struct.unpack(endian + 'I', data[e + 8:e + 12])[0]
                    if tag == 0x0132 and typ == 2 and cnt < 32:
                        out['DateTime'] = data[i + val_off:i + val_off + cnt].rstrip(b'\x00').decode('ascii', 'replace')
                    if tag == 0x0131 and typ == 2 and cnt < 128:
                        out['Software'] = data[i + val_off:i + val_off + cnt].rstrip(b'\x00').decode('ascii', 'replace')
                    if tag in (0x010F, 0x0110) and typ == 2 and 4 < cnt < 128:
                        out['Make' if tag == 0x010F else 'Model'] = data[i + val_off:i + val_off + cnt].rstrip(b'\x00').decode('ascii', 'replace')
                    if tag == 0x8769:
                        sn = struct.unpack(endian + 'H', data[i + val_off:i + val_off + 2])[0]
                        for kk in range(sn):
                            se = i + val_off + 2 + kk * 12
                            stag, styp, scnt = struct.unpack(endian + 'HHI', data[se:se + 8])
                            svo = struct.unpack(endian + 'I', data[se + 8:se + 12])[0]
                            if stag == 0x9003 and styp == 2 and scnt < 32:
                                out['DateTimeOriginal'] = data[i + svo:i + svo + scnt].rstrip(b'\x00').decode('ascii', 'replace')
                            if stag == 0x9004 and styp == 2 and scnt < 32:
                                out['DateTimeDigitized'] = data[i + svo:i + svo + scnt].rstrip(b'\x00').decode('ascii', 'replace')
                if out:
                    return out
            except Exception:
                continue
    return out


def _parse_exif_dt(s: str) -> Optional[datetime.datetime]:
    if not s:
        return None
    try:
        return datetime.datetime.strptime(s.strip(), "%Y:%m:%d %H:%M:%S")
    except (ValueError, AttributeError):
        return None


def _utc_to_local(dt: datetime.datetime) -> datetime.datetime:
    """Convert a naive UTC datetime to naive local time."""
    return dt.replace(tzinfo=datetime.timezone.utc).astimezone().replace(tzinfo=None)


def _mp4_header_atoms(path: str) -> Optional[bytes]:
    """Return the bytes of a video's top-level moov, meta and udta atoms,
    which hold its dates and device tags, skipping over the media data so a
    large video is not read in full. None if no moov atom can be reached."""
    out = b''
    found_moov = False
    with open(path, 'rb') as f:
        total = os.fstat(f.fileno()).st_size
        pos = 0
        while pos + 8 <= total:
            f.seek(pos)
            hdr = f.read(8)
            if len(hdr) < 8:
                break
            size, kind = struct.unpack('>I4s', hdr)
            hlen = 8
            if size == 1:  # 64-bit size follows the type
                ext = f.read(8)
                if len(ext) < 8:
                    break
                size = struct.unpack('>Q', ext)[0]
                hlen = 16
            elif size == 0:  # atom runs to the end of the file
                size = total - pos
            if size < hlen:
                break
            if kind in (b'moov', b'meta', b'udta'):
                out += hdr + f.read(size - hlen)
                found_moov = found_moov or kind == b'moov'
            pos += size
    return out if found_moov else None


def _mp4_times(path: str) -> dict:
    """Return {creation_time, modification_time} from the mvhd atom, plus
    {camera} when the file carries a device-make tag. The atom stores UTC;
    the times are returned in local time, which is what Windows Explorer
    shows as "Media created"."""
    out = {}
    data = _mp4_header_atoms(path)
    if data is None:
        # No moov atom found by walking the file: search all of it
        with open(path, 'rb') as f:
            data = f.read()
    for marker in (b'com.apple.quicktime.make', b'\xa9mak', b'com.android.manufacturer'):
        if marker in data:
            out['camera'] = marker.decode('latin-1')
            break
    i = data.find(b'mvhd')
    if i < 0:
        return out
    p = data[i + 4:]
    try:
        if p[0] == 1:
            ct = struct.unpack('>Q', p[4:12])[0]
            mt = struct.unpack('>Q', p[12:20])[0]
        else:
            ct = struct.unpack('>I', p[4:8])[0]
            mt = struct.unpack('>I', p[8:12])[0]
        epoch = datetime.datetime(1904, 1, 1)
        for key, v in (('creation_time', ct), ('modification_time', mt)):
            dt = epoch + datetime.timedelta(seconds=v)
            if dt.year > 1970:
                out[key] = _utc_to_local(dt)
    except Exception:
        pass
    return out


def _pil_meta(path: str) -> dict:
    """Return {DateTimeOriginal, DateTime, Software, Make, Model} from an
    image via Pillow."""
    out = {}
    if not HAS_PIL:
        return out
    try:
        with Image.open(path) as img:
            exif = img.getexif()
            v = exif.get_ifd(0x8769).get(0x9003)
            if v:
                out['DateTimeOriginal'] = v
            v = exif.get(0x0132)
            if v:
                out['DateTime'] = v
            v = exif.get(0x0131)
            if v:
                out['Software'] = v
            v = exif.get(0x010F)
            if v:
                out['Make'] = v
            v = exif.get(0x0110)
            if v:
                out['Model'] = v
    except Exception:
        pass
    return out


def _camera(d: dict) -> Optional[str]:
    """Join EXIF Make and Model into one string, or None if both are blank."""
    parts = [str(d.get(k) or '').strip('\x00 ') for k in ('Make', 'Model')]
    return ' '.join(p for p in parts if p) or None


def read_metadata(path: str, ftype: str) -> dict:
    """Return a normalized dict of {date_taken, modify_time, software, camera}."""
    out = {'date_taken': None, 'modify_time': None, 'software': None, 'camera': None}
    if ftype == '.heic':
        d = _heic_dates(path)
        out['date_taken'] = _parse_exif_dt(d.get('DateTimeOriginal') or d.get('DateTimeDigitized'))
        out['modify_time'] = _parse_exif_dt(d.get('DateTime'))
        out['software'] = d.get('Software')
        out['camera'] = _camera(d)
    elif ftype in VIDEO_EXTS:
        d = _mp4_times(path)
        out['date_taken'] = d.get('creation_time')
        out['modify_time'] = d.get('modification_time')
        out['camera'] = d.get('camera')
    elif ftype in ('.png', '.jpg', '.jpeg', '.tif', '.tiff'):
        d = _pil_meta(path)
        if not d and ftype in ('.tif', '.tiff'):
            # TIFF variants Pillow cannot open: walk the TIFF structure directly
            d = _heic_dates(path)
        out['date_taken'] = _parse_exif_dt(d.get('DateTimeOriginal'))
        out['modify_time'] = _parse_exif_dt(d.get('DateTime'))
        out['software'] = d.get('Software')
        out['camera'] = _camera(d)
    elif ftype == '.webp':
        out['camera'] = _camera(_pil_meta(path))
    return out


# ---------------- Core workflow ----------------

VIDEO_EXTS = ('.mov', '.mp4', '.m4v', '.3gp', '.3g2')

MEDIA_EXTS = {'.png', '.jpg', '.jpeg', '.heic', '.heif', *VIDEO_EXTS,
              '.tif', '.tiff', '.gif', '.bmp', '.webp'}

# Folder inside root that --quarantine-ads moves ad images into. Every phase
# ignores it.
QUARANTINE_DIR = '_ads'

# Folder inside root that --move-incomplete moves empty, null-filled and
# headerless files into. Every phase ignores it.
INCOMPLETE_DIR = '_incomplete'

# What INCOMPLETE_DIR was called before it also took headerless videos.
# --move-incomplete folds it into INCOMPLETE_DIR; every phase ignores it.
LEGACY_NO_DATA_DIR = '_no-image-data'

HOLDING_DIRS = (QUARANTINE_DIR, INCOMPLETE_DIR, LEGACY_NO_DATA_DIR)


def _prune_holding(root: str, dirpath: str, dirs: list) -> None:
    """Stop an os.walk of root from descending into the holding folders."""
    if dirpath == root:
        dirs[:] = [d for d in dirs if d not in HOLDING_DIRS]


# AVCHD camcorder video, and the older MPEG, Windows Media and AVI video that
# cameras and phones saved before MP4. Not in MEDIA_EXTS: their dates cannot
# be read or written here, so only --flatten handles them, going by the
# filename date.
CAMCORDER_EXTS = ('.m2ts', '.mts', '.mpg', '.mpeg', '.wmv', '.avi')

# Index files Sony's import software writes next to a file, named
# <file name>.modd / <file name>.moff. --flatten moves them with their file.
SIDECAR_EXTS = ('.modd', '.moff')


def iter_media(root: str, exts=None):
    """Yield (folder_name, filename, full_path) for every media file under
    root, or every file with one of exts if given."""
    exts = MEDIA_EXTS if exts is None else exts
    for folder in sorted(os.listdir(root)):
        fp = os.path.join(root, folder)
        if not os.path.isdir(fp) or folder in HOLDING_DIRS:
            continue
        for name in sorted(os.listdir(fp)):
            full = os.path.join(fp, name)
            if not os.path.isfile(full):
                continue
            ext = pathlib.Path(name).suffix.lower()
            if ext not in exts:
                continue
            yield folder, name, full


# YYYY-MM-DD or YYYYMMDD, optionally followed by HHMMSS, not embedded in a
# longer run of digits (so "Snapchat-1234567890" does not match).
_NAME_DATE_RE = re.compile(
    r'(?<!\d)(\d{4})([-_.]?)(\d{2})\2(\d{2})'
    r'(?:[-_ T.]?(\d{2})[-_.]?(\d{2})[-_.]?(\d{2})(?![0-9A-Za-z]))?(?!\d)')


def _filename_date(stem: str) -> Optional[datetime.datetime]:
    """Return the first plausible date embedded in a filename stem, at the
    embedded time if there is one, else 12:00:00."""
    today = datetime.date.today()
    for m in _NAME_DATE_RE.finditer(stem):
        try:
            d = datetime.date(int(m[1]), int(m[3]), int(m[4]))
        except ValueError:
            continue
        if d.year < 1990 or d > today:
            continue
        if m[5]:
            try:
                return datetime.datetime.combine(d, datetime.time(int(m[5]), int(m[6]), int(m[7])))
            except ValueError:
                pass
        return datetime.datetime.combine(d, datetime.time(12, 0, 0))
    return None


def _fallback_date(name: str, folder: str) -> tuple:
    """Return (datetime, source) to use when Date Taken is blank: the date in
    the filename, else the YYYY-MM-DD folder name at 12:00:00. The file's
    modified time is deliberately not consulted — imports reset it."""
    dt = _filename_date(pathlib.Path(name).stem)
    if dt is not None:
        return dt, 'filename'
    m = re.match(r'^(\d{4})-(\d{2})-(\d{2})$', folder)
    if m:
        return datetime.datetime(int(m[1]), int(m[2]), int(m[3]), 12, 0, 0), 'folder'
    return None, None


# M-D-YYYY folder names, e.g. "11-11-2019" or "1-30-2018"
_FOLDER_MDY_RE = re.compile(r'(?<!\d)(\d{1,2})[-_.](\d{1,2})[-_.](\d{4})(?!\d)')


def _folder_name_date(folder: str) -> Optional[datetime.datetime]:
    """Return the full date in a folder name at 12:00:00, or None. Accepts
    the forms _filename_date does (2019-08-30, 20190830) plus M-D-YYYY."""
    dt = _filename_date(folder)
    if dt is not None:
        return dt.replace(hour=12, minute=0, second=0)
    for m in _FOLDER_MDY_RE.finditer(folder):
        try:
            d = datetime.date(int(m[3]), int(m[1]), int(m[2]))
        except ValueError:
            continue
        if 1990 <= d.year and d <= datetime.date.today():
            return datetime.datetime.combine(d, datetime.time(12, 0, 0))
    return None


def _video_fallback_date(name: str, folders: list) -> tuple:
    """Return (datetime, source) to use when a video's Media Created is
    blank. folders are the names of the folders containing the file, nearest
    first. Priority: the date in the filename; else the nearest folder whose
    name holds a full date. A folder name holding only a year ("2014",
    "2013 West") is not a date and is never used."""
    dt = _filename_date(pathlib.Path(name).stem)
    if dt is not None:
        return dt, 'filename'
    for folder in folders:
        dt = _folder_name_date(folder)
        if dt is not None:
            return dt, 'folder'
    return None, None


def scan(root: str) -> list:
    """Return a list of plan dicts for every file under root."""
    rows = []
    for folder, name, full in iter_media(root):
        m = re.match(r'^(\d{4})-(\d{2})-(\d{2})$', folder)
        folder_date = datetime.date(int(m[1]), int(m[2]), int(m[3])) if m else None
        ftype = detect_type(full) or pathlib.Path(name).suffix.lower()
        meta = read_metadata(full, ftype)
        dt = meta['date_taken']
        # The name is kept as it is: a _YYYY-MM-DD ending is never stripped,
        # and the extension is Fix-Extensions.ps1's business
        p = pathlib.Path(name)
        fallback_source = _fallback_date(name, folder)[1] if dt is None else None
        if fallback_source == 'filename':
            action = 'fill-blank'
            target_folder = folder
            new_name = name
        elif fallback_source == 'folder':
            action = 'fill-blank'
            target_folder = folder
            # No date inside the file: the folder's date goes into the name,
            # once, so it survives the file leaving the folder
            dated = f"_{folder_date.isoformat()}"
            new_name = name if p.stem.endswith(dated) else f"{p.stem}{dated}{p.suffix}"
        elif dt is None:
            action = 'skip-no-folder-date'
            target_folder = folder
            new_name = name
        elif folder_date is None or dt.date() == folder_date:
            action = 'keep'
            target_folder = folder
            new_name = name
        else:
            action = 'move'
            target_folder = dt.date().isoformat()
            new_name = name
        rows.append({
            'current_folder': folder,
            'current_name': name,
            'detected_type': (ftype or '').lstrip('.').upper(),
            'date_taken': dt.strftime('%Y-%m-%d %H:%M:%S') if dt else '',
            'target_folder': target_folder,
            'new_name': new_name,
            'action': action,
        })
    return rows


# ---------------- Apply moves ----------------

def apply_moves(root: str, plan: list) -> dict:
    """Execute moves and renames. A 'fill-blank' row is not moved (its date
    is for --fill-blanks to write), but one still undated here gets the
    _<folder-date> suffix the plan gave it."""
    counts = Counter()
    for r in plan:
        src = os.path.join(root, r['current_folder'], r['current_name'])
        if r['action'] == 'fill-blank' and r['new_name'] != r['current_name']:
            dst = os.path.join(root, r['current_folder'], r['new_name'])
            if not os.path.isfile(src):
                counts['missing'] += 1
            elif os.path.exists(dst):
                counts['collision'] += 1
                print(f"COLLISION, skipping: {dst}", file=sys.stderr)
            else:
                shutil.move(src, dst)
                counts['suffixed-undated'] += 1
                print(f"RENAMED (no date in file): {r['current_folder']}/{r['current_name']} -> {r['new_name']}")
            continue
        if r['action'] not in ('move', 'keep'):
            counts['skipped-' + r['action']] += 1
            continue
        if not os.path.isfile(src):
            counts['missing'] += 1
            continue
        tgt_dir = os.path.join(root, r['target_folder'])
        os.makedirs(tgt_dir, exist_ok=True)
        dst = os.path.join(tgt_dir, r['new_name'])
        if src == dst:
            counts['unchanged'] += 1
            continue
        if os.path.exists(dst):
            counts['collision'] += 1
            print(f"COLLISION, skipping: {dst}", file=sys.stderr)
            continue
        shutil.move(src, dst)
        counts[r['action']] += 1
    return dict(counts)


# ---------------- Compare collisions ----------------

# The _<date-taken> suffix --flatten gives a file whose name is taken in root
_FLATTEN_SUFFIX_RE = re.compile(r'^(.+)_\d{4}-\d{2}-\d{2}$')


def _compare_pair(root: str, src: str, dst: str, counts: Counter) -> None:
    """Compare two colliding files byte for byte and report the result."""
    pair = f"{os.path.relpath(src, root)} | {os.path.relpath(dst, root)}"
    try:
        src_size, dst_size = os.path.getsize(src), os.path.getsize(dst)
        same = src_size == dst_size and filecmp.cmp(src, dst, shallow=False)
    except OSError as e:
        counts['failed'] += 1
        print(f"FAILED {pair}: {e}", file=sys.stderr)
        return
    if same:
        counts['identical'] += 1
        print(f"IDENTICAL ({src_size:,} bytes): {pair}")
        return
    dates = []
    for path in (src, dst):
        dt = read_metadata(path, detect_type(path) or pathlib.Path(path).suffix.lower())['date_taken']
        dates.append(dt.strftime('%Y-%m-%d %H:%M:%S') if dt else 'blank')
    taken = f"taken {dates[0]} vs {dates[1]}"
    if src_size == dst_size:
        counts['different-content'] += 1
        print(f"DIFFERENT (same size {src_size:,} bytes, content differs; {taken}): {pair}")
    else:
        counts['different-size'] += 1
        print(f"DIFFERENT ({src_size:,} vs {dst_size:,} bytes; {taken}): {pair}")


def _collision_pairs(root: str, plan: list):
    """Yield (file, file holding its name) for the three kinds of collision:
    a file --apply-moves skips because its name is taken in the folder for
    its real date; a file --flatten left in its subfolder because its name
    is taken in root; and a file --flatten renamed with a _<date-taken>
    suffix because its name was taken in root."""
    for r in plan:
        if r['action'] not in ('move', 'keep'):
            continue
        src = os.path.join(root, r['current_folder'], r['current_name'])
        dst = os.path.join(root, r['target_folder'], r['new_name'])
        if src == dst or not os.path.isfile(src) or not os.path.exists(dst) or os.path.samefile(src, dst):
            continue
        yield src, dst

    # Media files in root, grouped by name with any flatten suffix removed,
    # lowercased (Windows is case-insensitive)
    exts = MEDIA_EXTS | set(CAMCORDER_EXTS)
    plain, suffixed = {}, {}
    for name in sorted(os.listdir(root)):
        p = pathlib.Path(name)
        if p.suffix.lower() not in exts or not os.path.isfile(os.path.join(root, name)):
            continue
        m = _FLATTEN_SUFFIX_RE.match(p.stem)
        if m:
            suffixed.setdefault((m[1] + p.suffix).lower(), []).append(name)
        plain[name.lower()] = name
    for key, names in suffixed.items():
        if key in plain:
            for name in names:
                yield os.path.join(root, name), os.path.join(root, plain[key])
    for _folder, name, full in iter_media(root, exts):
        key = name.lower()
        for peer in ([plain[key]] if key in plain else []) + suffixed.get(key, []):
            yield full, os.path.join(root, peer)


def compare_collisions(root: str, plan: list) -> dict:
    """Report whether colliding files (see _collision_pairs) are duplicates,
    comparing each pair byte for byte. Changes nothing."""
    counts = Counter()
    for src, dst in _collision_pairs(root, plan):
        _compare_pair(root, src, dst, counts)
    return dict(counts)


# ---------------- Visual duplicate check ----------------

# Frames taken from each video, as fractions of its length
VISUAL_FRAME_POSITIONS = (0.1, 0.35, 0.6, 0.85)
# Both pictures are shrunk to a square this many pixels wide before their
# pixels are compared, which hides compression noise but not content
VISUAL_COMPARE_EDGE = 128
# Most that two pictures' width/height ratios, or two videos' lengths, may
# differ and still be the same picture or recording
VISUAL_SHAPE_TOLERANCE = 0.03
# Pixel difference (percent of full scale, averaged over the picture) up to
# which a pair gets each verdict, tightest first; anything above the last is
# DIFFERENT with high confidence. Measured on real files: a re-saved or
# re-compressed copy scores under 0.5, video frames a tenth of a second apart
# 0.6 to 5, the same scene seconds apart 4 to 20, unrelated pictures over 18.
VISUAL_VERDICTS = ((1.0, 'duplicate', 'high'), (2.0, 'duplicate', 'medium'),
                   (10.0, 'unsure', 'low'), (15.0, 'different', 'medium'))


def _format_pairs(root: str):
    """Yield pairs of media files in one folder (root or a subfolder) that
    share a name and differ only in extension, e.g. IMG_1.jpg and IMG_1.png:
    the files the cleanup script treats as duplicates."""
    folders = [''] + sorted(d for d in os.listdir(root)
                            if os.path.isdir(os.path.join(root, d)) and d not in HOLDING_DIRS)
    for folder in folders:
        fp = os.path.join(root, folder)
        groups = {}
        for name in sorted(os.listdir(fp)):
            p = pathlib.Path(name)
            if p.suffix.lower() in MEDIA_EXTS and os.path.isfile(os.path.join(fp, name)):
                groups.setdefault(p.stem.lower(), []).append(name)
        for names in groups.values():
            for other in names[1:]:
                yield os.path.join(fp, other), os.path.join(fp, names[0])


def _find_ffmpeg() -> Optional[str]:
    exe = shutil.which('ffmpeg')
    if exe:
        return exe
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:
        return None


def _video_frames(ffmpeg: str, path: str) -> tuple:
    """Return (length in seconds, [PIL images]) for frames taken at
    VISUAL_FRAME_POSITIONS through the video."""
    probe = subprocess.run([ffmpeg, '-hide_banner', '-i', path], capture_output=True, text=True, errors='replace')
    m = re.search(r'Duration: (\d+):(\d+):(\d+(?:\.\d+)?)', probe.stderr)
    if not m:
        raise ValueError('ffmpeg could not read the video length')
    length = int(m[1]) * 3600 + int(m[2]) * 60 + float(m[3])
    frames = []
    for pos in VISUAL_FRAME_POSITIONS:
        out = subprocess.run(
            [ffmpeg, '-hide_banner', '-loglevel', 'error', '-ss', f"{length * pos:.3f}", '-i', path,
             '-frames:v', '1', '-f', 'image2pipe', '-vcodec', 'mjpeg', '-q:v', '2', '-'],
            capture_output=True)
        if not out.stdout:
            raise ValueError(f"ffmpeg could not extract the frame at {length * pos:.1f}s")
        frames.append(Image.open(io.BytesIO(out.stdout)))
    return length, frames


def _visual_pictures(ffmpeg: Optional[str], path: str) -> tuple:
    """Return (video length or None, [PIL images]) to compare a file by: the
    picture itself for a photo, sampled frames for a video."""
    ftype = detect_type(path) or pathlib.Path(path).suffix.lower()
    if ftype in VIDEO_EXTS:
        if ffmpeg is None:
            raise ValueError('ffmpeg is needed to compare videos (pip install imageio-ffmpeg)')
        return _video_frames(ffmpeg, path)
    with Image.open(path) as img:
        return None, [img.convert('RGB')]


def _pixel_difference(a, b) -> float:
    """Return how far apart two pictures are, 0 (same) to 100 (opposite):
    the mean difference of their pixels once both are shrunk to one size."""
    size = (VISUAL_COMPARE_EDGE, VISUAL_COMPARE_EDGE)
    diff = ImageChops.difference(a.convert('RGB').resize(size), b.convert('RGB').resize(size))
    return sum(ImageStat.Stat(diff).mean) / 3 / 255 * 100


def _visual_verdict(ffmpeg: Optional[str], src: str, dst: str) -> tuple:
    """Return (verdict, confidence, reason) for whether two files show the
    same picture or recording."""
    src_len, src_pics = _visual_pictures(ffmpeg, src)
    dst_len, dst_pics = _visual_pictures(ffmpeg, dst)
    if (src_len is None) != (dst_len is None):
        return 'different', 'high', 'one is a photo, the other a video'
    if src_len is not None and abs(src_len - dst_len) > max(0.2, VISUAL_SHAPE_TOLERANCE * max(src_len, dst_len)):
        return 'different', 'high', f"video lengths differ, {src_len:.1f}s vs {dst_len:.1f}s"
    a, b = src_pics[0], dst_pics[0]
    if abs(a.width / a.height - b.width / b.height) > VISUAL_SHAPE_TOLERANCE * (a.width / a.height):
        return 'different', 'high', f"shapes differ, {a.width}x{a.height} vs {b.width}x{b.height}"
    # Judge a video by its least alike pair of frames
    worst = max(_pixel_difference(x, y) for x, y in zip(src_pics, dst_pics))
    what = f"pixel difference {worst:.1f}%"
    if src_len is not None:
        what += f" at the least alike of {len(src_pics)} frames, both {src_len:.1f}s long"
    if (a.width, a.height) != (b.width, b.height):
        what += f", sizes {a.width}x{a.height} vs {b.width}x{b.height}"
    for limit, verdict, confidence in VISUAL_VERDICTS:
        if worst <= limit:
            return verdict, confidence, what
    return 'different', 'high', what


def visual_compare(root: str, plan: list) -> dict:
    """Dry run only: judge whether pairs of files are duplicates by comparing
    their pixels, and report a verdict with a confidence level. Changes
    nothing and sends nothing anywhere. Covers every collision pair (see
    _collision_pairs) and every pair of same-named files in different formats
    (see _format_pairs). Byte-identical pairs need no pixel comparison."""
    counts = Counter()
    if not HAS_PIL:
        print("FAILED: Pillow is needed to compare pictures (pip install pillow)", file=sys.stderr)
        return {'failed': 1}
    ffmpeg = _find_ffmpeg()
    seen = set()
    for src, dst in list(_collision_pairs(root, plan)) + list(_format_pairs(root)):
        key = frozenset((os.path.normcase(src), os.path.normcase(dst)))
        if key in seen:
            continue
        seen.add(key)
        pair = f"{os.path.relpath(src, root)} | {os.path.relpath(dst, root)}"
        try:
            if os.path.getsize(src) == os.path.getsize(dst) and filecmp.cmp(src, dst, shallow=False):
                counts['identical'] += 1
                print(f"IDENTICAL (confidence certain, byte for byte): {pair}")
                continue
            verdict, confidence, reason = _visual_verdict(ffmpeg, src, dst)
        except Exception as e:
            counts['failed'] += 1
            print(f"FAILED {pair}: {e}", file=sys.stderr)
            continue
        counts[f"{verdict}-{confidence}"] += 1
        print(f"{verdict.upper()} (confidence {confidence}): {pair} - {reason}")
    return dict(counts)


# ---------------- Fill blank Date Taken ----------------

def _jpeg_set_exif(path: str, exif_bytes: bytes) -> None:
    """Replace (or insert) the EXIF APP1 segment of a JPEG, leaving the
    compressed image data byte-for-byte untouched."""
    if len(exif_bytes) > 65533:
        raise ValueError("EXIF block too large for a JPEG APP1 segment")
    with open(path, 'rb') as f:
        data = f.read()
    if data[:2] != b'\xff\xd8':
        raise ValueError("not a JPEG file")
    pos = insert_at = 2
    old = []
    while pos + 4 <= len(data) and data[pos] == 0xFF:
        marker = data[pos + 1]
        if marker in (0xDA, 0xD9):  # start of scan / end of image
            break
        end = pos + 2 + struct.unpack('>H', data[pos + 2:pos + 4])[0]
        if marker == 0xE0 and pos == insert_at:
            insert_at = end  # keep JFIF header(s) first
        elif marker == 0xE1 and data[pos + 4:pos + 10] == b'Exif\x00\x00':
            old.append((pos, end))
        pos = end
    segment = b'\xff\xe1' + struct.pack('>H', len(exif_bytes) + 2) + exif_bytes
    out = bytearray(data[:insert_at]) + segment
    cur = insert_at
    for s, e in old:
        out += data[cur:s]
        cur = e
    out += data[cur:]
    tmp = path + '.tmp'
    with open(tmp, 'wb') as f:
        f.write(out)
    os.replace(tmp, path)


def _write_exif_date_image(path: str, dt: datetime.datetime) -> None:
    """Write EXIF DateTime/DateTimeOriginal/DateTimeDigitized on JPEG or PNG.
    JPEGs get the EXIF segment swapped in place (no re-encode, no quality
    loss); other formats are re-saved through Pillow."""
    if not HAS_PIL:
        raise RuntimeError("Pillow is required to write EXIF; pip install pillow")
    ds = dt.strftime("%Y:%m:%d %H:%M:%S")
    tmp = path + '.tmp'
    with Image.open(path) as img:
        fmt = (img.format or '').upper()
        exif = img.getexif()
        exif[0x0132] = ds
        sub = exif.get_ifd(0x8769)
        sub[0x9003] = ds
        sub[0x9004] = ds
        exif_bytes = exif.tobytes()
        if fmt not in ('JPEG', 'MPO'):
            img.load()
            save_kwargs = {'exif': exif_bytes}
            if 'icc_profile' in img.info:
                save_kwargs['icc_profile'] = img.info['icc_profile']
            img.save(tmp, format=fmt, **save_kwargs)
    if fmt in ('JPEG', 'MPO'):
        _jpeg_set_exif(path, exif_bytes)
    else:
        os.replace(tmp, path)


def _find_exiftool() -> Optional[str]:
    """Return the path to exiftool, or None if it is not installed."""
    exe = shutil.which('exiftool')
    if exe:
        return exe
    # winget's default location, for terminals opened before it was on PATH
    local = os.environ.get('LOCALAPPDATA')
    if local:
        exe = os.path.join(local, 'Programs', 'ExifTool', 'ExifTool.exe')
        if os.path.isfile(exe):
            return exe
    return None


def _write_date_exiftool(exiftool: str, path: str, ftype: str, dt: datetime.datetime) -> None:
    """Write Date Taken with exiftool, for formats Pillow cannot write
    (HEIC, video). The file is rewritten in place; image data is not
    re-encoded."""
    ds = dt.strftime("%Y:%m:%d %H:%M:%S")
    if ftype in VIDEO_EXTS:
        # QuickTimeUTC: dt is local time, the container stores UTC
        tags = ['-api', 'QuickTimeUTC=1',
                '-QuickTime:CreateDate=' + ds, '-QuickTime:ModifyDate=' + ds,
                '-TrackCreateDate=' + ds, '-MediaCreateDate=' + ds]
    else:
        tags = ['-AllDates=' + ds]
    r = subprocess.run(
        [exiftool, '-charset', 'filename=utf8', '-overwrite_original', '-m', *tags, path],
        capture_output=True, text=True, encoding='utf-8', errors='replace')
    if r.returncode != 0 or '1 image files updated' not in r.stdout:
        msg = (r.stderr or r.stdout).strip().splitlines()
        raise RuntimeError(msg[-1] if msg else f"exiftool exited {r.returncode}")


def _iter_fill_candidates(root: str):
    """Yield (label, name, full_path, ftype, folders) for every file
    --fill-blanks considers: photos one level below root, like the other
    phases, and videos in root and every subfolder at any depth. folders are
    the names of the containing folders, nearest first, ending with root's
    own name."""
    root_name = os.path.basename(os.path.abspath(root))
    for dirpath, dirs, files in os.walk(root):
        _prune_holding(root, dirpath, dirs)
        dirs.sort()
        rel = os.path.relpath(dirpath, root)
        parts = [] if rel == '.' else rel.split(os.sep)
        for name in sorted(files):
            ext = pathlib.Path(name).suffix.lower()
            if ext not in MEDIA_EXTS:
                continue
            if len(parts) != 1 and ext not in VIDEO_EXTS:
                # Deeper files are taken at their extension's word, to avoid
                # opening every photo in a large tree
                continue
            full = os.path.join(dirpath, name)
            ftype = detect_type(full) or ext
            if len(parts) != 1 and ftype not in VIDEO_EXTS:
                continue
            yield '/'.join(parts + [name]), name, full, ftype, parts[::-1] + [root_name]


def fill_blanks(root: str, dry_run: bool = False) -> dict:
    counts = Counter()
    exiftool = _find_exiftool()
    verb = 'WOULD WRITE' if dry_run else 'WROTE'
    for label, name, full, ftype, folders in _iter_fill_candidates(root):
        if os.path.getsize(full) == 0:
            # A failed copy or transfer: there is no image to write a date into
            counts['empty-file'] += 1
            continue
        try:
            meta = read_metadata(full, ftype)
        except (OSError, MemoryError) as e:
            counts['read-failed'] += 1
            print(f"FAILED to read {label}: {e or type(e).__name__}", file=sys.stderr)
            continue
        if meta['date_taken'] is not None:
            counts['already-dated'] += 1
            continue
        if ftype in VIDEO_EXTS:
            field = 'Media Created'
            fill_dt, source = _video_fallback_date(name, folders)
        else:
            field = 'Date Taken'
            fill_dt, source = _fallback_date(name, folders[0])
        if fill_dt is None:
            counts['no-date-source'] += 1
            if ftype in VIDEO_EXTS:
                print(f"SKIPPED (no date in filename or folder names): {label}", file=sys.stderr)
            continue
        wrote = f"{verb} {field} {fill_dt:%Y-%m-%d %H:%M:%S} (from {source}"
        if ftype in ('.jpg', '.jpeg', '.png', '.tif', '.tiff'):
            try:
                if not dry_run:
                    try:
                        _write_exif_date_image(full, fill_dt)
                    except Exception:
                        # e.g. a TIFF variant Pillow cannot open
                        if not exiftool:
                            raise
                        _write_date_exiftool(exiftool, full, ftype, fill_dt)
                counts['written-from-' + source] += 1
                print(f"{wrote}): {label}")
            except Exception as e:
                counts['write-failed'] += 1
                print(f"FAILED to write {label}: {e}", file=sys.stderr)
        elif ftype in ('.heic', '.heif') + VIDEO_EXTS:
            if not exiftool:
                counts['heic-mov-needs-exiftool'] += 1
                print(f"SKIPPED (needs exiftool): {label}", file=sys.stderr)
                continue
            try:
                if not dry_run:
                    _write_date_exiftool(exiftool, full, ftype, fill_dt)
                counts['written-from-' + source] += 1
                print(f"{wrote}, exiftool): {label}")
            except Exception as e:
                counts['write-failed'] += 1
                print(f"FAILED to write {label}: {e}", file=sys.stderr)
        else:
            # GIF, BMP, WEBP: no Date Taken field this script can write
            counts['unsupported-format'] += 1
    if counts['empty-file']:
        print(f"NOTE: {counts['empty-file']} file(s) are empty (0 bytes) and were skipped.", file=sys.stderr)
    return dict(counts)


# ---------------- Convert PNG to JPEG ----------------

def _has_transparency(img) -> bool:
    """True when an image has at least one pixel that is not fully opaque.
    An alpha channel that is opaque throughout does not count."""
    if img.mode == 'P' and 'transparency' in img.info:
        img = img.convert('RGBA')
    if img.mode not in ('RGBA', 'LA', 'PA'):
        return 'transparency' in img.info
    return img.getchannel('A').getextrema()[0] < 255


def convert_pngs(root: str) -> dict:
    counts = Counter()
    if not HAS_PIL:
        print("Pillow required for PNG conversion; pip install pillow", file=sys.stderr)
        return {}
    for dirpath, dirs, files in os.walk(root):
        _prune_holding(root, dirpath, dirs)
        for name in files:
            if not name.lower().endswith('.png'):
                continue
            if pathlib.Path(name).stem.lower().endswith('-overlay'):
                # Snapchat caption layer: transparent, flattens to a blank image
                counts['skipped-overlay'] += 1
                continue
            p = os.path.join(dirpath, name)
            new = str(pathlib.Path(p).with_suffix('.jpg'))
            if os.path.exists(new):
                counts['already-converted'] += 1
                continue
            try:
                with Image.open(p) as img:
                    img.load()
                    if _has_transparency(img):
                        # JPEG has no transparency; flattening would change the picture
                        counts['skipped-transparent'] += 1
                        continue
                    exif_bytes = img.info.get('exif', b'')
                    img.convert('RGB').save(new, 'JPEG', quality=95, exif=exif_bytes, optimize=True)
            except (OSError, SyntaxError, ValueError, MemoryError) as e:
                # Not a picture despite the name (a macOS "._" sidecar, a
                # truncated download): report it and carry on with the rest
                counts['read-failed'] += 1
                print(f"FAILED to convert {os.path.relpath(p, root)}: {e or type(e).__name__}", file=sys.stderr)
                if os.path.exists(new):
                    # A half-written JPEG would pass for a finished conversion
                    os.remove(new)
                continue
            counts['converted'] += 1
    return dict(counts)


# ---------------- Sync timestamps (safe-only) ----------------

EDITOR_SOFTWARE_HINTS = ('photoshop', 'lightroom', 'camera raw', 'gimp',
                         'affinity', 'luminar', 'pixelmator', 'snapseed',
                         'vsco', 'instagram')


def _is_unedited(meta: dict) -> bool:
    """Heuristic for 'has NOT been manually modified since capture.'
    True when:
      - the EXIF file-level DateTime equals DateTimeOriginal (iOS updates
        DateTime when the user edits a photo), OR DateTime is missing;
      - the Software tag does not match a known editor.
    """
    sw = (meta.get('software') or '').lower()
    for hint in EDITOR_SOFTWARE_HINTS:
        if hint in sw:
            return False
    dt = meta.get('date_taken')
    mt = meta.get('modify_time')
    if dt is None:
        return False
    if mt is None:
        return True
    # Equal within one minute of tolerance (iOS may differ by seconds)
    return abs((mt - dt).total_seconds()) < 60


def _set_creation_time(path: str, ts: float) -> bool:
    """Set the filesystem "Date Created" to the POSIX timestamp ts.
    os.utime only covers accessed/modified, so this goes through the Win32
    SetFileTime API. Returns False on platforms where it isn't supported."""
    if os.name != 'nt':
        return False
    import ctypes
    from ctypes import wintypes
    k32 = ctypes.WinDLL('kernel32', use_last_error=True)
    k32.CreateFileW.restype = wintypes.HANDLE
    k32.CreateFileW.argtypes = (wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD,
                                wintypes.LPVOID, wintypes.DWORD, wintypes.DWORD,
                                wintypes.HANDLE)
    k32.SetFileTime.restype = wintypes.BOOL
    k32.SetFileTime.argtypes = (wintypes.HANDLE, ctypes.POINTER(wintypes.FILETIME),
                                ctypes.POINTER(wintypes.FILETIME),
                                ctypes.POINTER(wintypes.FILETIME))
    k32.CloseHandle.argtypes = (wintypes.HANDLE,)
    # FILETIME = 100ns ticks since 1601-01-01 UTC
    ticks = int(round((ts + 11644473600) * 10_000_000))
    ft = wintypes.FILETIME(ticks & 0xFFFFFFFF, ticks >> 32)
    # FILE_WRITE_ATTRIBUTES, share read/write/delete, OPEN_EXISTING
    handle = k32.CreateFileW(os.path.abspath(path), 0x0100, 0x7, None, 3, 0x80, None)
    if handle == wintypes.HANDLE(-1).value:
        raise ctypes.WinError(ctypes.get_last_error())
    try:
        if not k32.SetFileTime(handle, ctypes.byref(ft), None, None):
            raise ctypes.WinError(ctypes.get_last_error())
    finally:
        k32.CloseHandle(handle)
    return True


def sync_timestamps(root: str) -> dict:
    """Set Date Created to Date Taken on every dated file, and Date Modified
    too on unedited ones. An edited file keeps its Date Modified, which is
    when it was edited."""
    counts = Counter()
    for folder, name, full in iter_media(root):
        ftype = detect_type(full) or pathlib.Path(name).suffix.lower()
        meta = read_metadata(full, ftype)
        if meta['date_taken'] is None:
            counts['no-date-taken'] += 1
            continue
        ts = meta['date_taken'].timestamp()
        if not _is_unedited(meta):
            try:
                if _set_creation_time(full, ts):
                    counts['created-only-edited'] += 1
                    print(f"SYNCED created only (edited) {meta['date_taken']:%Y-%m-%d %H:%M:%S}: {folder}/{name}")
                else:
                    counts['skipped-edited'] += 1
                    print(f"SKIP (edited): {folder}/{name}")
            except Exception as e:
                counts['failed'] += 1
                print(f"FAILED {folder}/{name}: {e}", file=sys.stderr)
            continue
        try:
            os.utime(full, (ts, ts))
            if not _set_creation_time(full, ts):
                counts['created-unsupported'] += 1
            counts['synced'] += 1
            print(f"SYNCED {meta['date_taken']:%Y-%m-%d %H:%M:%S}: {folder}/{name}")
        except Exception as e:
            counts['failed'] += 1
            print(f"FAILED {folder}/{name}: {e}", file=sys.stderr)
    return dict(counts)


# ---------------- Flatten into root ----------------

def flatten(root: str, dry_run: bool = False) -> dict:
    """Move dated files up from their subfolder into root, then remove the
    subfolders that end up empty. Undated files are left in place, except
    inside the _ads folder: ads rarely have a Date Taken, so there every file
    moves up from its subfolder into _ads itself, dated or not."""
    counts = Counter()
    verb = 'WOULD MOVE' if dry_run else 'MOVED'
    gone = set()  # paths moved or removed, so a dry run can tell what empties
    ads = os.path.join(root, QUARANTINE_DIR)
    for base in (root, ads):
        if not os.path.isdir(base):
            continue
        in_ads = base == ads
        prefix = QUARANTINE_DIR + '/' if in_ads else ''
        # Names already claimed in base, lowercased (Windows is case-insensitive)
        taken = {n.lower() for n in os.listdir(base)}
        if not in_ads:
            root_names = taken
        exts = MEDIA_EXTS if in_ads else MEDIA_EXTS | set(CAMCORDER_EXTS)
        for folder, name, full in iter_media(base, exts):
            from_folder = False
            if in_ads:
                # The subfolder name is the only date an ad has
                suffix = folder
            elif pathlib.Path(name).suffix.lower() in CAMCORDER_EXTS:
                dt = _filename_date(pathlib.Path(name).stem)
                if dt is None:
                    # The folder's date goes into the name, so it survives
                    # the file leaving the folder
                    dt = _folder_name_date(folder)
                    from_folder = dt is not None
                if dt is None:
                    counts['left-undated'] += 1
                    print(f"LEFT (no date in filename or folder name): {folder}/{name}")
                    continue
                if _has_no_data(full):
                    counts['left-no-data'] += 1
                    print(f"LEFT (no video data): {folder}/{name}", file=sys.stderr)
                    continue
                suffix = f"{dt:%Y-%m-%d}"
            else:
                ftype = detect_type(full) or pathlib.Path(name).suffix.lower()
                dt = read_metadata(full, ftype)['date_taken']
                if dt is None:
                    counts['left-undated'] += 1
                    print(f"LEFT (no Date Taken): {folder}/{name}")
                    continue
                suffix = f"{dt:%Y-%m-%d}"
            new_name = name
            if from_folder or new_name.lower() in taken:
                p = pathlib.Path(name)
                new_name = f"{p.stem}_{suffix}{p.suffix}"
                if new_name.lower() in taken:
                    counts['collision'] += 1
                    print(f"COLLISION, skipping: {prefix}{folder}/{name}", file=sys.stderr)
                    continue
                counts['dated-from-folder' if from_folder else 'renamed'] += 1
            if not dry_run:
                try:
                    shutil.move(full, os.path.join(base, new_name))
                except OSError as e:
                    counts['failed'] += 1
                    print(f"FAILED {prefix}{folder}/{name}: {e}", file=sys.stderr)
                    continue
            taken.add(new_name.lower())
            gone.add(full)
            counts['ads-moved' if in_ads else 'moved'] += 1
            print(f"{verb}: {prefix}{folder}/{name} -> {prefix}{new_name}")
            for ext in SIDECAR_EXTS:
                # Sidecars follow their file, taking its new name if it changed
                side = full + ext
                if not os.path.isfile(side):
                    continue
                side_dst = os.path.join(base, new_name + ext)
                if (new_name + ext).lower() in taken:
                    counts['collision'] += 1
                    print(f"COLLISION, skipping: {prefix}{folder}/{name}{ext}", file=sys.stderr)
                    continue
                if not dry_run:
                    try:
                        shutil.move(side, side_dst)
                    except OSError as e:
                        counts['failed'] += 1
                        print(f"FAILED {prefix}{folder}/{name}{ext}: {e}", file=sys.stderr)
                        continue
                taken.add((new_name + ext).lower())
                gone.add(side)
                counts['sidecars-moved'] += 1

    # Videos in a photo folder (or photos in a video folder) are easy to lose
    # sight of once everything is in one place: say how many of the less
    # common kind root itself holds after the moves
    kinds = Counter()
    for n in root_names:
        ext = pathlib.Path(n).suffix
        if ext in VIDEO_EXTS + CAMCORDER_EXTS:
            kinds['video'] += 1
        elif ext in MEDIA_EXTS:
            kinds['photo'] += 1
    if kinds:
        major, minor = ('photo', 'video') if kinds['photo'] >= kinds['video'] else ('video', 'photo')
        counts[f'{minor}s-in-{major}-folder'] = kinds[minor]
        print(f"OPPOSITE TYPE: {kinds[minor]:,} {minor}(s) in this {major} folder "
              f"({kinds[major]:,} {major}(s)){' once flattened' if dry_run else ''}")

    # A sidecar is only any use beside the file it describes: delete the ones
    # in root and its subfolders whose file is not there
    verb = 'WOULD DELETE' if dry_run else 'DELETED'
    subfolders = sorted(d for d in os.listdir(root)
                        if os.path.isdir(os.path.join(root, d)) and d not in HOLDING_DIRS)
    for folder in [''] + subfolders:
        fp = os.path.join(root, folder)
        for name in sorted(os.listdir(fp)):
            p = pathlib.Path(name)
            full = os.path.join(fp, name)
            if p.suffix.lower() not in SIDECAR_EXTS or full in gone or not os.path.isfile(full):
                continue
            if os.path.exists(os.path.join(fp, p.stem)):
                continue
            rel = f"{folder}/{name}" if folder else name
            if not dry_run:
                try:
                    os.remove(full)
                except OSError as e:
                    counts['failed'] += 1
                    print(f"FAILED {rel}: {e}", file=sys.stderr)
                    continue
            gone.add(full)
            counts['orphan-sidecars-deleted'] += 1
            print(f"{verb} (sidecar with no file): {rel}")

    verb = 'WOULD REMOVE EMPTY' if dry_run else 'REMOVED EMPTY'
    # Incomplete files stay in their dated folders: the name is their only date
    kept = tuple(os.path.join(root, d) + os.sep for d in (INCOMPLETE_DIR, LEGACY_NO_DATA_DIR))
    for dirpath, _dirs, _files in os.walk(root, topdown=False):
        if dirpath in (root, ads) or (dirpath + os.sep).startswith(kept):
            continue
        if any(os.path.join(dirpath, e) not in gone for e in os.listdir(dirpath)):
            counts['folders-kept'] += 1
            continue
        if not dry_run:
            try:
                os.rmdir(dirpath)
            except OSError as e:
                counts['folders-kept'] += 1
                print(f"FAILED to remove {dirpath}: {e}", file=sys.stderr)
                continue
        gone.add(dirpath)
        counts['folders-removed'] += 1
        print(f"{verb}: {dirpath}")
    return dict(counts)


# ---------------- Opposite type report ----------------

# A year folder: "2014", "2013 West"
_YEAR_FOLDER_RE = re.compile(r'^(19|20)\d{2}(?!\d)')


def _year_folders(root: str) -> list:
    """Return (label, path) for each year folder in root, or root itself as
    the one folder if it holds none."""
    years = sorted(d for d in os.listdir(root)
                   if _YEAR_FOLDER_RE.match(d) and os.path.isdir(os.path.join(root, d)))
    return [(y, os.path.join(root, y)) for y in years] or [(os.path.basename(os.path.normpath(root)), root)]


def _media_by_kind(root: str, tops: list) -> tuple:
    """Return (found, major, minor): found maps each year folder to
    {'photo': [paths relative to root], 'video': [...]}, searched at every
    depth with holding folders left out; major is the kind there is more of
    in all, minor the other."""
    found = {}
    for label, top in tops:
        kinds = found[label] = {'photo': [], 'video': []}
        for dirpath, dirs, files in os.walk(top):
            dirs[:] = sorted(d for d in dirs if d not in HOLDING_DIRS)
            for name in sorted(files):
                ext = pathlib.Path(name).suffix.lower()
                if ext in VIDEO_EXTS + CAMCORDER_EXTS:
                    kind = 'video'
                elif ext in MEDIA_EXTS:
                    kind = 'photo'
                else:
                    continue
                rel = os.path.relpath(os.path.join(dirpath, name), root)
                kinds[kind].append(rel.replace(os.sep, '/'))
    total = {k: sum(len(kinds[k]) for kinds in found.values()) for k in ('photo', 'video')}
    major, minor = ('photo', 'video') if total['photo'] >= total['video'] else ('video', 'photo')
    return found, major, minor


def opposite_type(root: str) -> dict:
    """Report photos filed among videos, or videos among photos, one year
    folder at a time. Whichever kind root holds more of decides what it is.
    Each year folder is searched at every depth, holding folders excepted.
    With no year folders in root, root itself is reported as one folder.
    Nothing is changed."""
    tops = _year_folders(root)
    found, major, minor = _media_by_kind(root, tops)
    total = {k: sum(len(kinds[k]) for kinds in found.values()) for k in ('photo', 'video')}
    counts = {f'{minor}s-in-{major}-folders': total[minor],
              'year-folders': len(tops),
              f'year-folders-with-{minor}s': sum(1 for kinds in found.values() if kinds[minor])}
    for label, kinds in found.items():
        for rel in kinds[minor]:
            print(f"{minor.upper()} IN {major.upper()} FOLDER: {rel}")
    for label, kinds in found.items():
        counts[label] = len(kinds[minor])
        print(f"{label}: {len(kinds[minor]):,} {minor}(s), {len(kinds[major]):,} {major}(s)")
    print(f"OPPOSITE TYPE: {total[minor]:,} {minor}(s) among {total[major]:,} {major}(s), in "
          f"{counts[f'year-folders-with-{minor}s']} of {len(tops)} year folder(s)")
    return counts


# ---------------- Final check report ----------------

# An audit workbook, whatever was put before or after "media_audit"
_AUDIT_FILE_RE = re.compile(r'media_audit.*\.xlsx$', re.IGNORECASE)

def _count_files(folder: str, skip_holding: bool = False) -> int:
    """Count the files in folder at every depth."""
    n = 0
    for _dirpath, dirs, files in os.walk(folder):
        if skip_holding:
            dirs[:] = [d for d in dirs if d not in HOLDING_DIRS]
        n += len(files)
    return n


def _media_type(folder: str) -> str:
    """Return 'photos' or 'videos', whichever kind folder holds more of at
    any depth (holding folders left out), or '' if it holds neither."""
    n = Counter()
    for _dirpath, dirs, files in os.walk(folder):
        dirs[:] = [d for d in dirs if d not in HOLDING_DIRS]
        for name in files:
            ext = pathlib.Path(name).suffix.lower()
            if ext in VIDEO_EXTS + CAMCORDER_EXTS:
                n['videos'] += 1
            elif ext in MEDIA_EXTS:
                n['photos'] += 1
    if not n:
        return ''
    return 'photos' if n['photos'] >= n['videos'] else 'videos'


def _audit_prefix(folder: str) -> str:
    """Return the prefix an audit workbook in folder carries,
    "<folder name>_<photos|videos>_" (2014_videos_). A drive root has no
    name and a folder with no media no type; either part is then left out."""
    name = os.path.basename(os.path.normpath(os.path.abspath(folder)))
    return ''.join(f"{part}_" for part in (name, _media_type(folder)) if part)


def final_check(root: str, own_xlsx: Optional[str] = None, dry_run: bool = False) -> dict:
    """Report what the other phases left behind, one year folder at a time:
    files still in _ads, files still in _incomplete, photos among videos (or
    videos among photos), subfolders --flatten did not remove, and audit
    workbooks. own_xlsx is the workbook this run is recorded in, which is
    not reported. With no year folders in root, root itself is reported as
    one folder. The only change made: an audit workbook whose name does not
    start with its folder's name and media type is renamed so that it does."""
    tops = _year_folders(root)
    found, major, minor = _media_by_kind(root, tops)
    opposite = f'{minor}s-in-{major}-folders'
    counts = Counter({'ads-files': 0, 'incomplete-files': 0, opposite: 0,
                      'folders-not-flattened': 0, 'collisions': 0, 'collisions-identical': 0,
                      'collisions-different': 0, 'audit-files': 0,
                      'year-folders': len(tops), 'year-folders-with-oddities': 0})

    def is_audit(folder: str, name: str) -> bool:
        full = os.path.join(folder, name)
        return (bool(_AUDIT_FILE_RE.search(name)) and os.path.isfile(full)
                and not (own_xlsx and os.path.exists(own_xlsx) and os.path.samefile(full, own_xlsx)))

    def report_audit(folder: str, name: str, tally: Counter) -> None:
        """List one audit workbook, giving it its folder's name and media
        type as a prefix (2014_videos_) if it has none, so workbooks from
        different folders can be told apart once they are gathered in one
        place."""
        tally['audit-files'] += 1
        rel = os.path.relpath(os.path.join(folder, name), root).replace(os.sep, '/')
        new_name = _audit_prefix(folder) + name[name.lower().index('media_audit'):]
        if new_name.lower() == name.lower():
            print(f"AUDIT FILE: {rel}")
            return
        if os.path.exists(os.path.join(folder, new_name)):
            tally['audit-files-not-renamed'] += 1
            print(f"AUDIT FILE: {rel}")
            print(f"COLLISION, not renaming: {rel} ({new_name} exists)", file=sys.stderr)
            return
        if not dry_run:
            try:
                os.rename(os.path.join(folder, name), os.path.join(folder, new_name))
            except OSError as e:
                tally['failed'] += 1
                print(f"AUDIT FILE: {rel}")
                print(f"FAILED to rename {rel}: {e}", file=sys.stderr)
                return
        tally['audit-files-renamed'] += 1
        print(f"{'WOULD RENAME' if dry_run else 'RENAMED'} AUDIT FILE: {rel} -> {new_name}")

    if tops[0][1] != root:
        # Workbooks from runs on root itself sit beside the year folders
        for name in sorted(os.listdir(root)):
            if is_audit(root, name):
                report_audit(root, name, counts)
    summary = []
    for label, top in tops:
        year = Counter()
        for dirpath, dirs, files in os.walk(top):
            for name in sorted(files):
                if is_audit(dirpath, name):
                    report_audit(dirpath, name, year)
            for d in sorted(dirs):
                if d not in HOLDING_DIRS:
                    continue
                n = _count_files(os.path.join(dirpath, d))
                if not n:
                    continue
                key = 'ads-files' if d == QUARANTINE_DIR else 'incomplete-files'
                year[key] += n
                rel = os.path.relpath(os.path.join(dirpath, d), root).replace(os.sep, '/')
                print(f"{'IN ADS' if d == QUARANTINE_DIR else 'INCOMPLETE'}: {rel} ({n:,} file(s))")
            dirs[:] = [d for d in dirs if d not in HOLDING_DIRS]
        for rel in found[label][minor]:
            print(f"{minor.upper()} IN {major.upper()} FOLDER: {rel}")
        year[opposite] = len(found[label][minor])
        # Name clashes --flatten met: no plan is given, so the ones only
        # --apply-moves would meet (which need every file's date) are left out
        for src, dst in _collision_pairs(top, []):
            how = 'renamed with a date' if os.path.dirname(src) == top else 'left in its subfolder'
            try:
                same = os.path.getsize(src) == os.path.getsize(dst) and filecmp.cmp(src, dst, shallow=False)
            except OSError as e:
                year['failed'] += 1
                print(f"FAILED to compare {os.path.relpath(src, root)}: {e}", file=sys.stderr)
                continue
            verdict = 'identical' if same else 'different'
            year['collisions'] += 1
            year[f'collisions-{verdict}'] += 1
            pair = ' | '.join(os.path.relpath(p, root).replace(os.sep, '/') for p in (src, dst))
            print(f"COLLISION, {verdict} files ({how}): {pair}")
        for d in sorted(os.listdir(top)):
            sub = os.path.join(top, d)
            if not os.path.isdir(sub) or d in HOLDING_DIRS:
                continue
            # With no year folders, root's own subfolders are named as they are
            rel = d if top == root else f"{label}/{d}"
            year['folders-not-flattened'] += 1
            print(f"NOT FLATTENED: {rel} ({_count_files(sub, skip_holding=True):,} file(s))")
        counts.update(year)
        if any(year.values()):
            counts['year-folders-with-oddities'] += 1
        summary.append(
            f"{label}: {year['ads-files']:,} in _ads, {year['incomplete-files']:,} in _incomplete, "
            f"{year[opposite]:,} {minor}(s) among {major}s, "
            f"{year['folders-not-flattened']:,} folder(s) not flattened, "
            f"{year['collisions']:,} collision(s), "
            f"{year['audit-files']:,} audit file(s)" if any(year.values())
            else f"{label}: clean")
    for line in summary:
        print(line)
    print(f"FINAL CHECK: {counts['ads-files']:,} in _ads, {counts['incomplete-files']:,} in _incomplete, "
          f"{counts[opposite]:,} {minor}(s) among {major}s, "
          f"{counts['folders-not-flattened']:,} folder(s) not flattened, "
          f"{counts['collisions']:,} collision(s) "
          f"({counts['collisions-identical']:,} identical, {counts['collisions-different']:,} different), "
          f"{counts['audit-files']:,} audit file(s); "
          f"{counts['year-folders-with-oddities']} of {len(tops)} year folder(s) have something left")
    return dict(counts)


# ---------------- Quarantine ad images and videos ----------------

# Filename stems that ad images and videos are saved under:
#   0c4cda27-b4cc-4e92-a446-d6b780f24a64   UUID
#   bed51b94a_1595                         9-char hex id (with a letter, so
#                                          20231220_142355 is safe) + number
#   gmsnet2                                gmsnet + optional number
#   news_images%2F1714658759372            news_images + URL-encoded "/" + number
#   UnityAdsCache-3d4c8e33...b18d2091      UnityAdsCache- + 64-char hex hash
#   54ac2fda0a6755305200011c-b30-600       24-char hex id + one or more
#   55d58970f6cd4574f635d6e3_568-1443241239
#                                          -/_ separated letter-digit parts
#   Update_Now_Video_V2_720x1280_15s___fudxlv
#                                          pixel dimensions somewhere in the
#                                          name, ending ___ + 6-char id
# optionally followed by the _YYYY-MM-DD suffix --flatten adds on a name
# clash, or a " (2)" copy marker.
_AD_NAME_RE = re.compile(
    r'(?:[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}'
    r'|(?=\d*[a-f])[0-9a-f]{9}_\d+'
    r'|gmsnet\d*'
    r'|news_images%2F\d+'
    r'|UnityAdsCache-[0-9a-f]{64}'
    r'|(?=\d*[a-f])[0-9a-f]{24}(?:[-_][0-9a-z]+)+'
    r'|.*(?<!\d)\d{3,4}x\d{3,4}(?!\d).*___[0-9a-z]{6})'
    r'(?:_\d{4}-\d{2}-\d{2})?(?: \(\d+\))?',
    re.IGNORECASE)


# Stems that mark a video (only) as an ad, with the same optional suffixes:
#   32129eda9b8e718c5277, vuyyzy0brvod5gcaocnv, agqnsutqkbwulksalejd
#       20 lowercase letters and digits with at least one letter. Lowercase
#       only, so camera names such as VID20231220142355123 are safe.
#   313e81d7...daa11f.mp4-720x1280-h264-Q2, peacock_..._rev-720x1280-Q2
#       ends in -<width>x<height>-Q2 or -<width>x<height>-h264-Q2, the
#       transcode tag ad networks append
#   d5a210d011c47b8300bb8034048719a9
#       32-character lowercase hex hash
_AD_VIDEO_NAME_RE = re.compile(
    r'(?:(?=\d*[a-z])[0-9a-z]{20}'
    r'|[0-9a-f]{32}'
    r'|.*-\d{3,4}x\d{3,4}(?:-h264)?-Q2)'
    r'(?:_\d{4}-\d{2}-\d{2})?(?: \(\d+\))?')


def is_ad_name(name: str) -> bool:
    p = pathlib.Path(name)
    if _AD_NAME_RE.fullmatch(p.stem):
        return True
    return p.suffix.lower() in VIDEO_EXTS and _AD_VIDEO_NAME_RE.fullmatch(p.stem) is not None


def quarantine_ads(root: str, dry_run: bool = False) -> dict:
    """Move ad images and videos from root and every subfolder into root/_ads,
    keeping their relative path. A file is an ad only if its name matches AND
    it has no camera info (EXIF Make/Model for a photo, a device-make tag for
    a video); name matches with camera info are kept and listed.
    Emptied folders are left for --flatten or the cleanup script."""
    counts = Counter()
    if not HAS_PIL:
        # Without Pillow the camera check cannot run, and name alone is not enough
        print("Pillow required to check for camera info; pip install pillow", file=sys.stderr)
        return {}
    verb = 'WOULD QUARANTINE' if dry_run else 'QUARANTINED'
    for dirpath, dirs, files in os.walk(root):
        _prune_holding(root, dirpath, dirs)
        dirs.sort()
        for name in sorted(files):
            if pathlib.Path(name).suffix.lower() not in MEDIA_EXTS or not is_ad_name(name):
                continue
            full = os.path.join(dirpath, name)
            rel = os.path.relpath(full, root)
            ftype = detect_type(full) or pathlib.Path(name).suffix.lower()
            camera = read_metadata(full, ftype)['camera']
            if camera:
                counts['kept-camera-info'] += 1
                print(f"KEPT (camera: {camera}): {rel}")
                continue
            dst = os.path.join(root, QUARANTINE_DIR, rel)
            if os.path.exists(dst):
                counts['collision'] += 1
                print(f"COLLISION, skipping: {rel}", file=sys.stderr)
                continue
            if not dry_run:
                try:
                    os.makedirs(os.path.dirname(dst), exist_ok=True)
                    shutil.move(full, dst)
                except OSError as e:
                    counts['failed'] += 1
                    print(f"FAILED {rel}: {e}", file=sys.stderr)
                    continue
            counts['quarantined'] += 1
            print(f"{verb}: {rel}")
    return dict(counts)


# ---------------- Move incomplete files ----------------

def _has_no_data(path: str) -> bool:
    """True when a file is 0 bytes or contains nothing but null bytes."""
    with open(path, 'rb') as f:
        while True:
            chunk = f.read(1 << 20)
            if not chunk:
                return True
            if chunk.count(0) != len(chunk):
                return False


def _incomplete_kind(path: str) -> Optional[str]:
    """Return why a media file is incomplete, or None if it is not: 'empty'
    (0 bytes), 'null-filled' (nothing but null bytes), or 'no-video-header'
    (a video with no moov atom: a stub or a truncated copy, which no player
    can open)."""
    if os.path.getsize(path) == 0:
        return 'empty'
    if pathlib.Path(path).suffix.lower() in VIDEO_EXTS:
        if _mp4_header_atoms(path) is not None:
            return None
        if _has_no_data(path):
            return 'null-filled'
        # A photo wearing a video extension is misnamed, not incomplete
        return 'no-video-header' if detect_type(path) in (None,) + VIDEO_EXTS else None
    return 'null-filled' if _has_no_data(path) else None


def _fold_legacy_no_data(root: str, dry_run: bool, counts: Counter) -> None:
    """Move what an earlier version put in root/_no-image-data into
    root/_incomplete, keeping relative paths."""
    legacy = os.path.join(root, LEGACY_NO_DATA_DIR)
    if not os.path.isdir(legacy):
        return
    verb = 'WOULD MOVE' if dry_run else 'MOVED'
    for dirpath, dirs, files in os.walk(legacy, topdown=False):
        for name in sorted(files):
            full = os.path.join(dirpath, name)
            rel = os.path.relpath(full, legacy)
            dst = os.path.join(root, INCOMPLETE_DIR, rel)
            if os.path.exists(dst):
                counts['collision'] += 1
                print(f"COLLISION, skipping: {LEGACY_NO_DATA_DIR}{os.sep}{rel}", file=sys.stderr)
                continue
            if not dry_run:
                try:
                    os.makedirs(os.path.dirname(dst), exist_ok=True)
                    shutil.move(full, dst)
                except OSError as e:
                    counts['failed'] += 1
                    print(f"FAILED {LEGACY_NO_DATA_DIR}{os.sep}{rel}: {e}", file=sys.stderr)
                    continue
            counts['moved-from-' + LEGACY_NO_DATA_DIR] += 1
        if not dry_run and not os.listdir(dirpath):
            os.rmdir(dirpath)
    n = counts['moved-from-' + LEGACY_NO_DATA_DIR]
    if n:
        print(f"{verb}: {n} file(s) from {LEGACY_NO_DATA_DIR} into {INCOMPLETE_DIR}")


def move_incomplete(root: str, dry_run: bool = False) -> dict:
    """Move incomplete media files (see _incomplete_kind) from root and every
    subfolder into root/_incomplete, keeping their relative path (the dated
    folder name may be the only record of when they were taken). Emptied
    folders are left for --flatten or the cleanup script."""
    counts = Counter()
    verb = 'WOULD MOVE' if dry_run else 'MOVED'
    _fold_legacy_no_data(root, dry_run, counts)
    for dirpath, dirs, files in os.walk(root):
        _prune_holding(root, dirpath, dirs)
        dirs.sort()
        for name in sorted(files):
            if pathlib.Path(name).suffix.lower() not in MEDIA_EXTS:
                continue
            full = os.path.join(dirpath, name)
            rel = os.path.relpath(full, root)
            try:
                kind = _incomplete_kind(full)
            except OSError as e:
                counts['failed'] += 1
                print(f"FAILED {rel}: {e}", file=sys.stderr)
                continue
            if kind is None:
                continue
            dst = os.path.join(root, INCOMPLETE_DIR, rel)
            if os.path.exists(dst):
                counts['collision'] += 1
                print(f"COLLISION, skipping: {rel}", file=sys.stderr)
                continue
            if not dry_run:
                try:
                    os.makedirs(os.path.dirname(dst), exist_ok=True)
                    shutil.move(full, dst)
                except OSError as e:
                    counts['failed'] += 1
                    print(f"FAILED {rel}: {e}", file=sys.stderr)
                    continue
            counts['moved-' + kind] += 1
            print(f"{verb} ({kind}): {rel}")
    return dict(counts)


# ---------------- Delete junk files ----------------

# Thumbnail caches Windows leaves in picture folders (usually hidden). Whole
# names, not a pattern: "west sucking thumb.jpg" is a photo.
THUMBS_NAMES = ('thumbs.db', 'ehthumbs.db', 'ehthumbs_vista.db')

# First four bytes of an AppleDouble file, the "._name" sidecar macOS writes
# beside every file it copies to a non-Mac drive
APPLEDOUBLE_MAGIC = b'\x00\x05\x16\x07'


def _is_appledouble(full: str) -> bool:
    """True for a macOS "._" sidecar. The name alone is not trusted: the file
    must also open with the AppleDouble signature, so a picture that merely
    starts with "._" is kept."""
    if not os.path.basename(full).startswith('._'):
        return False
    try:
        with open(full, 'rb') as f:
            return f.read(4) == APPLEDOUBLE_MAGIC
    except OSError:
        return False


def delete_junk(root: str, dry_run: bool = False) -> dict:
    """Delete every .AAE file (iPhone edit sidecar), Windows thumbnail cache
    (Thumbs.db) and macOS "._" sidecar in root and every subfolder. Emptied
    folders are left for --flatten or the cleanup script."""
    counts = Counter()
    verb = 'WOULD DELETE' if dry_run else 'DELETED'
    for dirpath, dirs, files in os.walk(root):
        dirs.sort()
        for name in sorted(files):
            full = os.path.join(dirpath, name)
            # Checked first: "._IMG_0001.AAE" is a sidecar, not an AAE file
            if _is_appledouble(full):
                kind = 'deleted-mac-sidecars'
            elif pathlib.Path(name).suffix.lower() == '.aae':
                kind = 'deleted'
            elif name.lower() in THUMBS_NAMES:
                kind = 'deleted-thumbs'
            else:
                continue
            rel = os.path.relpath(full, root)
            if not dry_run:
                try:
                    try:
                        os.remove(full)
                    except PermissionError:
                        # Thumbs.db is often read-only as well as hidden
                        os.chmod(full, 0o666)
                        os.remove(full)
                except OSError as e:
                    counts['failed'] += 1
                    print(f"FAILED {rel}: {e}", file=sys.stderr)
                    continue
            counts[kind] += 1
            print(f"{verb}: {rel}")
    return dict(counts)


# ---------------- Reporting ----------------

def write_xlsx(root: str, plan: list, output: str) -> None:
    try:
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("openpyxl not installed; skipping XLSX report. (pip install openpyxl)", file=sys.stderr)
        return
    # Keep the run log sheets of an existing workbook; only the scan sheets
    # are rebuilt
    wb = _open_audit(output)
    for title in ('Summary', 'Plan', 'Only moves'):
        if title in wb.sheetnames:
            del wb[title]
    header_fill = PatternFill('solid', start_color='305496')
    header_font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    body_font = Font(name='Arial', size=10)
    thin = Side(border_style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    def style(ws, ncols, widths):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='left', vertical='center')
            cell.border = border
        for r in range(2, ws.max_row + 1):
            for c in range(1, ncols + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = body_font
                cell.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
                cell.border = border
        ws.freeze_panes = 'A2'
        ws.row_dimensions[1].height = 22
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    ws = wb.create_sheet('Summary', 0)
    ws.append(['Metric', 'Value'])
    ws.append(['Files scanned', len(plan)])
    ws.append(['Keep in place', sum(1 for r in plan if r['action'] == 'keep')])
    ws.append(['Move to date-matched folder', sum(1 for r in plan if r['action'] == 'move')])
    ws.append(['Blank Date Taken (would fill)', sum(1 for r in plan if r['action'] == 'fill-blank')])
    style(ws, 2, [44, 10])

    ws2 = wb.create_sheet('Plan', 1)
    cols = ['Current folder', 'Current name', 'Detected type', 'Date Taken',
            'Target folder', 'New name', 'Action']
    ws2.append(cols)
    for r in plan:
        ws2.append([r['current_folder'], r['current_name'], r['detected_type'],
                    r['date_taken'], r['target_folder'], r['new_name'], r['action']])
    style(ws2, len(cols), [14, 32, 10, 22, 14, 32, 24])

    ws3 = wb.create_sheet('Only moves', 2)
    ws3.append(['From folder', 'Current name', 'Date Taken', 'Target folder', 'New name'])
    for r in plan:
        if r['action'] == 'move':
            ws3.append([r['current_folder'], r['current_name'], r['date_taken'],
                        r['target_folder'], r['new_name']])
    style(ws3, 5, [14, 32, 22, 14, 32])

    _save_audit(wb, output)


# Sheets every run appends to, so the workbook is a history of what was done
LOG_SHEET = 'Run log'
COUNTS_SHEET = 'Run counts'

# A reported line: its action is the leading run of capitalised words
# ("MOVED", "WOULD WRITE", "FAILED"), the rest is the detail
_AUDIT_LINE_RE = re.compile(r'^([A-Z]{2,}(?: [A-Z]{2,})*)\b[ ,:]*(.*)$')
# The confidence level --visual-compare gives its verdicts, for its own column
_AUDIT_CONFIDENCE_RE = re.compile(r'^\(confidence (\w+)')
_XLSX_BAD_CHARS = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')


def _open_audit(path: str):
    """Load the audit workbook, or start an empty one if there is none."""
    from openpyxl import Workbook, load_workbook
    if os.path.isfile(path):
        try:
            return load_workbook(path)
        except Exception as e:
            print(f"Could not read {path} ({e}); starting a new audit workbook.", file=sys.stderr)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _save_audit(wb, path: str) -> bool:
    try:
        wb.save(path)
        return True
    except PermissionError:
        print(f"Could not write {path}: close it in Excel and run again. "
              "This run is NOT recorded in the audit workbook.", file=sys.stderr)
        return False


def append_audit_log(xlsx: str, started: datetime.datetime, runs: list) -> None:
    """Append what each phase reported to the audit workbook. runs is a list
    of (phase, dry_run, counts, lines); every line that starts with an action
    word becomes a row of the Run log sheet, every count a row of Run counts."""
    try:
        from openpyxl.styles import Font
    except ImportError:
        print("openpyxl not installed; run not recorded in the audit workbook. (pip install openpyxl)", file=sys.stderr)
        return
    wb = _open_audit(xlsx)

    def sheet(title, header, widths):
        if title in wb.sheetnames:
            return wb[title]
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font = Font(name='Arial', bold=True)
        ws.freeze_panes = 'A2'
        for col, w in zip('ABCDEF', widths):
            ws.column_dimensions[col].width = w
        return ws

    log = sheet(LOG_SHEET, ['Run started', 'Phase', 'Dry run', 'Action', 'Detail', 'Confidence'], [20, 18, 9, 22, 110, 12])
    if log.cell(row=1, column=6).value is None:
        # A workbook started before --visual-compare existed
        log.cell(row=1, column=6, value='Confidence').font = Font(name='Arial', bold=True)
        log.column_dimensions['F'].width = 12
    tally = sheet(COUNTS_SHEET, ['Run started', 'Phase', 'Dry run', 'Count', 'Value'], [20, 18, 9, 34, 10])
    when = started.strftime('%Y-%m-%d %H:%M:%S')
    for phase, dry_run, counts, lines in runs:
        dry = 'yes' if dry_run else 'no'
        for k, v in counts.items():
            tally.append([when, phase, dry, k, v])
        for ln in lines:
            m = _AUDIT_LINE_RE.match(ln.strip())
            if m:
                conf = _AUDIT_CONFIDENCE_RE.match(m[2])
                log.append([when, phase, dry, m[1], _XLSX_BAD_CHARS.sub('?', m[2]), conf[1] if conf else None])
    if _save_audit(wb, xlsx):
        print(f"Audit workbook: {xlsx}")


# ---------------- CLI ----------------

# Order phases run in when several are given on one command line: clear out
# junk first, then date, convert, move, sync, and flatten last (once files are
# in the root the other phases no longer see them).
PHASE_ORDER = ('delete_junk', 'move_incomplete', 'quarantine_ads', 'fill_blanks', 'convert_png',
               'apply_moves', 'compare_collisions', 'visual_compare', 'sync_timestamps', 'flatten',
               'opposite_type', 'final_check')
ALL_PHASES = ('delete_junk', 'move_incomplete', 'fill_blanks', 'convert_png', 'apply_moves', 'sync_timestamps')
DRY_RUN_PHASES = ('flatten', 'quarantine_ads', 'move_incomplete', 'delete_junk', 'fill_blanks', 'visual_compare',
                  'final_check')


def _default_xlsx(root: str) -> str:
    """Return today's audit workbook name for root,
    <root folder name>_<photos|videos>_media_audit_<date>.xlsx
    (2014_videos_media_audit_2026-10-07.xlsx). A workbook from earlier today
    under the old name, with no prefix, is renamed to it so the day's runs
    stay in one workbook."""
    plain = f'media_audit_{datetime.date.today():%Y-%m-%d}.xlsx'
    if not os.path.isdir(root):
        return plain
    # Today's workbook keeps the name it was started under, even if the
    # folder's media type has changed since
    folder = os.path.basename(os.path.normpath(os.path.abspath(root)))
    for name in sorted(os.listdir(root)):
        if folder and name.lower().startswith(f"{folder}_".lower()) and name.lower().endswith(f"_{plain}"):
            return name
    name = _audit_prefix(root) + plain
    old, new = os.path.join(root, plain), os.path.join(root, name)
    if name != plain and os.path.isfile(old) and not os.path.exists(new):
        try:
            os.rename(old, new)
        except OSError:
            pass  # open in Excel: it stays, and the final check reports it
    return name

# Most problem lines repeated per phase in the closing summary
MAX_SUMMARY_ISSUES = 50


class _IssueTee:
    """Pass a stream through unchanged while keeping a copy of each line, so
    the problems a phase reported can be repeated in the closing summary and
    everything it reported can be recorded in the audit workbook."""

    def __init__(self, stream, log=None):
        self.stream = stream
        self.text = ''
        # Shared list that completed lines are also appended to, in the
        # order written, for the audit workbook
        self.log = log
        self._partial = ''

    def write(self, s):
        self.text += s
        if self.log is not None:
            *done, self._partial = (self._partial + s).split('\n')
            self.log.extend(ln for ln in done if ln.strip())
        return self.stream.write(s)

    def flush(self):
        self.stream.flush()

    def lines(self) -> list:
        return [ln for ln in self.text.splitlines() if ln.strip()]


def _run_phase(phase: str, args) -> dict:
    if phase == 'delete_junk':
        return delete_junk(args.root, args.dry_run)
    if phase == 'move_incomplete':
        return move_incomplete(args.root, args.dry_run)
    if phase == 'quarantine_ads':
        return quarantine_ads(args.root, args.dry_run)
    if phase == 'fill_blanks':
        return fill_blanks(args.root, args.dry_run)
    if phase == 'convert_png':
        return convert_pngs(args.root)
    if phase == 'sync_timestamps':
        return sync_timestamps(args.root)
    if phase == 'flatten':
        return flatten(args.root, args.dry_run)
    if phase == 'opposite_type':
        return opposite_type(args.root)
    if phase == 'final_check':
        return final_check(args.root, os.path.join(args.root, args.xlsx), args.dry_run)
    if phase == 'compare_collisions':
        return compare_collisions(args.root, scan(args.root))
    if phase == 'visual_compare':
        return visual_compare(args.root, scan(args.root))
    # report-only and apply-moves both need the scanned plan
    plan = scan(args.root)
    out_xlsx = os.path.join(args.root, args.xlsx)
    write_xlsx(args.root, plan, out_xlsx)
    if phase == 'apply_moves':
        return apply_moves(args.root, plan)
    return dict(Counter(r['action'] for r in plan))


def _print_summary(results: list) -> None:
    """Print every phase's counts and reported problems together, so a run of
    several phases can be read from the end of the output."""
    print()
    print('=' * 24 + ' Summary of this run ' + '=' * 24)
    for key, counts, issues, _log in results:
        print(f"\n{key}")
        if not counts:
            print("    nothing to do")
        width = max((len(k) for k in counts), default=0)
        for k, v in counts.items():
            print(f"    {k:<{width}}  {v:>7}")
        if issues:
            print(f"    problems reported: {len(issues)}")
            for ln in issues[:MAX_SUMMARY_ISSUES]:
                print(f"      {ln}")
            if len(issues) > MAX_SUMMARY_ISSUES:
                print(f"      ... and {len(issues) - MAX_SUMMARY_ISSUES} more (see above)")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--root', required=True, help='Photo folder root (e.g. D:\\Pictures\\2026)')
    ap.add_argument('--report-only', action='store_true', help='Scan and write audit XLSX; no changes. Cannot be combined with other phases.')
    ap.add_argument('--all', action='store_true', help='Run --delete-junk, --move-incomplete, --fill-blanks, --convert-png, --apply-moves and --sync-timestamps in that order, then print one summary of all of them.')
    ap.add_argument('--fill-blanks', action='store_true', help='Write Date Taken into EXIF for files whose Date Taken is blank, and Media Created into videos (at any depth) whose Media Created is blank, using the date in the filename, else the folder date.')
    ap.add_argument('--convert-png', action='store_true', help='Convert PNG files to JPEG (preserving EXIF).')
    ap.add_argument('--apply-moves', action='store_true', help='Move files whose Date Taken differs from the folder name.')
    ap.add_argument('--compare-collisions', action='store_true', help='Report only: compare each file --apply-moves skips as a COLLISION with the file already holding its name in the target folder, and list the pair as IDENTICAL or DIFFERENT. Also covers files --flatten renamed with a _<date-taken> suffix or left in their subfolder. Changes nothing. Not part of --all.')
    ap.add_argument('--visual-compare', action='store_true', help='Dry run only (give --dry-run too): compare the pixels of each pair of possible duplicates (collision pairs, and same-named files in different formats) and report DUPLICATE, DIFFERENT or UNSURE with a confidence level. Runs entirely on this computer; videos need ffmpeg (pip install imageio-ffmpeg). Changes nothing. Not part of --all.')
    ap.add_argument('--sync-timestamps', action='store_true', help='Set Date Created to Date Taken, and Date Modified too for files that have not been edited.')
    ap.add_argument('--flatten', action='store_true', help='Move files that have a Date Taken up out of their subfolders into the root, then delete the emptied subfolders. Inside _ads, files move up into _ads itself even with no Date Taken.')
    ap.add_argument('--opposite-type', action='store_true', help='Report only: give the folder that holds the year folders (D:\\Videos) as --root. Lists every photo filed among the videos, or video among the photos, with a count for each year folder. Changes nothing. Not part of --all.')
    ap.add_argument('--final-check', action='store_true', help='Give the folder that holds the year folders (D:\\Videos) as --root. For each year folder, reports files still in _ads, files still in _incomplete, photos among videos (or videos among photos), subfolders not flattened and audit workbooks. Renames an audit workbook not yet named for its folder and media type; changes nothing else. Not part of --all.')
    ap.add_argument('--quarantine-ads', action='store_true', help='Move ad images and videos (ad-style filename AND no camera info) from the root and every subfolder into an _ads folder for review. Nothing is deleted.')
    ap.add_argument('--move-incomplete', '--move-no-data', action='store_true', help='Move incomplete media files (0 bytes, all null bytes, or a video with no video header) from the root and every subfolder into an _incomplete folder, keeping their subfolder path. Nothing is deleted.')
    ap.add_argument('--delete-junk', action='store_true', help='Permanently delete every .AAE file (iPhone edit sidecar), Windows thumbnail cache (Thumbs.db) and macOS "._" sidecar in the root and every subfolder.')
    ap.add_argument('--dry-run', action='store_true', help='With --fill-blanks, --flatten, --quarantine-ads, --move-incomplete, --delete-junk or --final-check: list what would be written, moved, renamed, removed or deleted without changing anything.')
    ap.add_argument('--xlsx', help='Audit workbook filename (relative to root). Defaults to <root folder name>_<photos|videos>_media_audit_<today>.xlsx (2014_videos_media_audit_2026-10-07.xlsx), so each folder and each day gets its own workbook. Every run appends what it did to the Run log and Run counts sheets.')
    ap.add_argument('--audit-log', metavar='NAME', help='Record lines read from standard input in the audit workbook as phase NAME, then exit. Used by Fix-Extensions.ps1 and Cleanup-Pictures.ps1.')
    args = ap.parse_args(argv)
    if args.xlsx is None:
        args.xlsx = _default_xlsx(args.root)
    started = datetime.datetime.now()
    out_xlsx = os.path.join(args.root, args.xlsx)

    if args.audit_log:
        if not os.path.isdir(args.root):
            print(f"Not a directory: {args.root}", file=sys.stderr)
            return 2
        lines = sys.stdin.buffer.read().decode('utf-8-sig', 'replace').splitlines()
        append_audit_log(out_xlsx, started, [(args.audit_log, args.dry_run, {}, lines)])
        return 0

    phases = [p for p in PHASE_ORDER
              if getattr(args, p) or (args.all and p in ALL_PHASES)]
    if args.report_only:
        if phases:
            ap.error('--report-only cannot be combined with other phases')
        phases = ['report_only']
    if not phases:
        ap.error('choose at least one phase (e.g. --report-only, --fill-blanks, --all)')
    if args.dry_run and any(p not in DRY_RUN_PHASES for p in phases):
        ap.error('--dry-run is only supported with --fill-blanks, --flatten, --quarantine-ads, --move-incomplete, --delete-junk, --visual-compare and --final-check')
    if args.visual_compare and not args.dry_run:
        ap.error('--visual-compare only reports for now: add --dry-run')

    if not os.path.isdir(args.root):
        print(f"Not a directory: {args.root}", file=sys.stderr)
        return 2

    results = []
    for phase in phases:
        key = phase + '_dry_run' if args.dry_run else phase
        if len(phases) > 1:
            print(f"\n----- {key} -----")
        log = []
        tee = _IssueTee(sys.stderr, log)
        out = _IssueTee(sys.stdout, log)
        sys.stderr, sys.stdout = tee, out
        try:
            counts = _run_phase(phase, args)
        finally:
            sys.stderr, sys.stdout = tee.stream, out.stream
        print(json.dumps({key: counts}, indent=2))
        results.append((key, counts, tee.lines(), log))
    if len(phases) > 1:
        _print_summary(results)
    append_audit_log(out_xlsx, started,
                     [(phase, args.dry_run, counts, log)
                      for phase, (_key, counts, _issues, log) in zip(phases, results)])
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
